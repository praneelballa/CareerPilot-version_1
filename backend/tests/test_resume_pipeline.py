import io
import pytest
from pathlib import Path
import docx
import fitz
from app.services.extractor import ResumeExtractorService
from app.services.skill_normalizer import SkillNormalizer
from app.schemas.candidate import SkillCategories


def test_skill_normalization():
    raw_skills = SkillCategories(
        programming_languages=["python3", "Python", "JS", "typescript", "golang"],
        frameworks=["reactjs", "FastAPI"],
        libraries=["PyTorch", "pytorch"],
        databases=["Postgres", "mongodb"],
        cloud_and_tools=["docker", "k8s", "aws"],
        ai_and_ml=["machine learning", "NLP", "LLM"],
        soft_skills=["Communication", "Leadership"]
    )
    normalized = SkillNormalizer.normalize_skills(raw_skills)
    
    assert normalized.programming_languages == ["Python", "JavaScript", "TypeScript", "Go"]
    assert "PostgreSQL" in normalized.databases
    assert "Docker" in normalized.cloud_and_tools
    assert "Kubernetes" in normalized.cloud_and_tools
    assert "LLMs" in normalized.ai_and_ml
    assert len(normalized.libraries) == 1  # Deduplicated PyTorch


def test_docx_deterministic_extraction(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Jane Doe", 0)
    doc.add_paragraph("Experienced Software Engineer skilled in Python and React.")
    doc.save(str(doc_path))

    extracted = ResumeExtractorService.extract_from_docx(doc_path)
    assert "Jane Doe" in extracted
    assert "Experienced Software Engineer" in extracted


def test_pdf_deterministic_extraction(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "John Smith - Distributed Systems Architect with Go and Kubernetes.")
    doc.save(str(pdf_path))
    doc.close()

    extracted = ResumeExtractorService.extract_from_pdf(pdf_path)
    assert "John Smith" in extracted
    assert "Distributed Systems Architect" in extracted


def test_empty_file_rejection(tmp_path: Path):
    empty_file = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(empty_file))
    doc.close()

    with pytest.raises(Exception):
        ResumeExtractorService.extract_text(empty_file, "pdf")


def test_upload_resume_endpoint(client):
    # Create test DOCX in memory
    doc = docx.Document()
    doc.add_heading("Alex Rivera", 0)
    doc.add_paragraph("Full Stack Developer with 4 years building FastAPI and React systems.")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = client.post(
        "/api/v1/resume/upload",
        files={"file": ("alex_rivera_resume.docx", buffer, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )

    assert response.status_code == 201
    data = response.json()
    assert "candidate_id" in data
    assert data["filename"] == "alex_rivera_resume.docx"
    assert "Alex Rivera" in data["extracted_text_preview"]