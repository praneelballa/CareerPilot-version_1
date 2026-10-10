import fitz  # PyMuPDF
import docx
from pathlib import Path
from fastapi import HTTPException, status


class ResumeExtractorService:
    @staticmethod
    def extract_from_pdf(file_path: Path) -> str:
        text_content = []
        try:
            with fitz.open(str(file_path)) as doc:
                for page_num, page in enumerate(doc):
                    page_text = page.get_text("text")
                    if page_text:
                        text_content.append(page_text)
            return "\n\n".join(text_content).strip()
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unable to parse PDF content: {str(e)}"
            )

    @staticmethod
    def extract_from_docx(file_path: Path) -> str:
        text_content = []
        try:
            doc = docx.Document(str(file_path))
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_content.append(" | ".join(row_text))
            return "\n".join(text_content).strip()
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unable to parse DOCX document: {str(e)}"
            )

    @classmethod
    def extract_text(cls, file_path: Path, extension: str) -> str:
        ext = extension.lower().strip(".")
        if ext == "pdf":
            raw_text = cls.extract_from_pdf(file_path)
        elif ext in ["docx", "doc"]:
            raw_text = cls.extract_from_docx(file_path)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file extension: .{ext}. Supported formats: PDF, DOCX"
            )

        if not raw_text or len(raw_text.strip()) < 20:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Extracted resume text is empty or too short to analyze."
            )
        return raw_text